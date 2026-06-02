"""简历文件解析服务 - 支持 PDF/DOCX 上传解析"""

import json
import re
from pathlib import Path
from typing import Any

from backend.config import settings


class ResumeParser:
    """解析上传的 PDF 或 DOCX 简历文件"""

    UPLOAD_DIR = Path(settings.DATA_DIR) / "uploads"

    def __init__(self):
        self.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

    async def save_upload(self, file_content: bytes, filename: str) -> Path:
        """保存上传文件到本地"""
        # 安全处理文件名
        safe_name = re.sub(r'[^\w\-. ]', '_', filename)
        file_path = self.UPLOAD_DIR / safe_name

        # 避免重名覆盖
        counter = 1
        while file_path.exists():
            stem = file_path.stem
            suffix = file_path.suffix
            file_path = self.UPLOAD_DIR / f"{stem}_{counter}{suffix}"
            counter += 1

        file_path.write_bytes(file_content)
        return file_path

    async def parse_pdf(self, file_path: Path) -> str:
        """解析 PDF 文件为纯文本"""
        try:
            import fitz
            doc = fitz.open(str(file_path))
            text_parts = []
            for page in doc:
                text_parts.append(page.get_text())
            doc.close()
            return "\n".join(text_parts)
        except ImportError:
            return "PDF 解析库未安装 (PyMuPDF)"
        except Exception as e:
            return f"PDF 解析失败: {str(e)}"

    async def parse_docx(self, file_path: Path) -> str:
        """解析 DOCX 文件为纯文本"""
        try:
            from docx import Document
            doc = Document(str(file_path))
            text_parts = []

            # 提取段落
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text.strip())

            # 提取表格
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)

            return "\n".join(text_parts)
        except ImportError:
            return "DOCX 解析库未安装 (python-docx)"
        except Exception as e:
            return f"DOCX 解析失败: {str(e)}"

    async def parse_file(self, file_path: Path) -> str:
        """根据扩展名解析文件"""
        ext = file_path.suffix.lower()
        if ext == ".pdf":
            return await self.parse_pdf(file_path)
        elif ext in (".docx", ".doc"):
            return await self.parse_docx(file_path)
        else:
            return f"不支持的文件格式: {ext}"

    async def extract_structured_data(self, text: str) -> dict[str, Any]:
        """从纯文本简历中提取结构化信息（基础解析，后续 AI 会优化）"""
        lines = text.strip().split("\n")
        lines = [l.strip() for l in lines if l.strip()]

        data: dict[str, Any] = {
            "name": "",
            "phone": "",
            "email": "",
            "title": "",
            "summary": "",
            "skills": [],
            "work_experiences": [],
            "projects": [],
            "educations": [],
            "raw_text": text[:5000],  # 保留原始文本供 AI 优化
        }

        # 简单正则提取联系方式
        email_pattern = r'[\w.+-]+@[\w-]+\.[\w.-]+'
        phone_pattern = r'(1[3-9]\d{9})'

        for line in lines:
            # 邮箱
            emails = re.findall(email_pattern, line)
            if emails and not data["email"]:
                data["email"] = emails[0]

            # 手机号
            phones = re.findall(phone_pattern, line)
            if phones and not data["phone"]:
                data["phone"] = phones[0]

        # 将第一行作为姓名（通常在简历顶部）
        if lines:
            data["name"] = lines[0]

        return data
